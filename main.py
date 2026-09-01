from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import List
from openpyxl.styles import PatternFill, Font
from contextlib import asynccontextmanager
import pandas as pd
import math
import json
import io
import httpx
import sqlite3
import traceback
import time

DB_FILE = "pi_director.db"
MARKET_CACHE = {"data": {}, "time": 0} # Кэш для рынка Житы

def init_db():
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS characters (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            character_id INTEGER UNIQUE,
            name TEXT NOT NULL,
            ccu_level INTEGER DEFAULT 0,
            ic_level INTEGER DEFAULT 0,
            access_token TEXT,
            refresh_token TEXT
        )
    ''')
    conn.commit()
    conn.close()

init_db()

def get_db_connection():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn

PI_TYPE_NAMES = [
    "Broadcast Node", "Integrity Response Drones", "Nano-Factory", 
    "Organic Mortar Applicators", "Recursive Computing Module", 
    "Self-Harmonizing Power Core", "Sterile Conduits", "Wetware Mainframe",
    "Biotech Research Reports", "Camera Drones", "Condensates", 
    "Cryoprotectant Solution", "Data Chips", "Gel-Matrix Biopaste", 
    "Guidance Systems", "Hazmat Detection Systems", "Hermetic Membranes", 
    "High-Tech Transmitters", "Industrial Explosives", "Neocoms", 
    "Nuclear Reactors", "Planetary Vehicles", "Robotics", 
    "Smartfab Units", "Supercomputers", "Synthetic Synapses", 
    "Transcranial Microcontrollers", "Ukomi Superconductors", "Vaccines",
    "Biocells", "Construction Blocks", "Consumer Electronics", "Coolant", 
    "Enriched Uranium", "Fertilizer", "Genetically Enhanced Livestock", 
    "Livestock", "Mechanical Parts", "Microfiber Shielding", "Miniature Electronics", 
    "Nanites", "Oxides", "Polyaramids", "Polytextiles", "Rocket Fuel", 
    "Silicate Glass", "Superconductors", "Supertensile Plastics", "Synthetic Oil", 
    "Test Cultures", "Transmitter", "Viral Agent", "Water-Cooled CPU",
    "Water", "Industrial Fibers", "Reactive Metals", "Biofuels", "Proteins", 
    "Silicon", "Toxic Metals", "Electrolytes", "Bacteria", "Oxygen", 
    "Precious Metals", "Chiral Structures", "Biomass", "Oxidizing Compound", "Plasmoids",
    "Aqueous Liquids", "Autotrophs", "Base Metals", "Carbon Compounds", "Complex Organisms", 
    "Felsic Magma", "Heavy Metals", "Ionic Solutions", "Microorganisms", "Noble Gas", 
    "Noble Metals", "Non-CS Crystals", "Planktic Colonies", "Reactive Gas", "Suspended Plasma",
    "Barren Command Center", "Gas Command Center", "Ice Command Center", 
    "Lava Command Center", "Oceanic Command Center", "Plasma Command Center", 
    "Storm Command Center", "Temperate Command Center"
]

PI_TYPE_IDS = {}

@asynccontextmanager
async def lifespan(app: FastAPI):
    try:
        async with httpx.AsyncClient() as client:
            response = await client.post(
                "https://esi.evetech.net/latest/universe/ids/",
                json=PI_TYPE_NAMES,
                headers={'Accept-Language': 'en'}
            )
            data = response.json()
            if 'inventory_types' in data:
                for item in data['inventory_types']:
                    PI_TYPE_IDS[item['name']] = item['id']
                print(f"[OK] Загружены TypeID для {len(PI_TYPE_IDS)} товаров.")
    except Exception as e:
        print("[ERROR] ESI:", e)
    yield 
    PI_TYPE_IDS.clear()

app = FastAPI(title="PI Director API", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware, allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"]
)

class SystemRequest(BaseModel): constellations: List[str]
class PlanRequest(BaseModel): constellations: List[str]; factory_sys: str; target_products: List[str]
class ExportRequest(BaseModel): plan_data: list

def load_planet_data():
    df = pd.read_csv('planet industry.csv', sep=';', skiprows=1, encoding='utf-8')
    df = df.dropna(subset=['Constellation'])
    
    # Принудительно переводим в строку и обрезаем скрытые пробелы по краям
    df['Constellation'] = df['Constellation'].astype(str).str.strip()
    
    df = df[df['Constellation'] != 'max P2']
    
    # Фильтруем любые варианты ошибок Excel (#REF!, #reference, #REFERENCE), игнорируя регистр
    df = df[~df['Constellation'].str.contains('#REF', case=False, na=False)]
    
    return df

async def get_market_prices():
    now = time.time()
    # Кэшируем цены на 1 час, чтобы не спамить Fuzzwork
    if now - MARKET_CACHE["time"] < 3600 and MARKET_CACHE["data"]:
        return MARKET_CACHE["data"]
    type_ids_str = ",".join(str(v) for v in PI_TYPE_IDS.values())
    try:
        async with httpx.AsyncClient(timeout=10.0) as client:
            response = await client.get(f"https://market.fuzzwork.co.uk/aggregates/?region=10000002&types={type_ids_str}")
            market_data = response.json()
        prices = {}
        for prod_name, type_id in PI_TYPE_IDS.items():
            str_id = str(type_id)
            if str_id in market_data:
                prices[prod_name] = float(market_data[str_id]["buy"]["max"])
        MARKET_CACHE["data"] = prices
        MARKET_CACHE["time"] = now
        return prices
    except: return MARKET_CACHE["data"]

@app.get("/api/auth/callback")
def auth_callback(code: str):
    conn = get_db_connection()
    cursor = conn.cursor()
    mock_chars_from_esi = [(90001, "Factory Chief 1", 5, 5, "t", "r"), (90002, "Factory Chief 2", 5, 5, "t", "r")] + [(90004 + i, f"Miner {i+1}", 4, 5, "t", "r") for i in range(10)]
    for char in mock_chars_from_esi:
        cursor.execute('''INSERT INTO characters (character_id, name, ccu_level, ic_level, access_token, refresh_token) VALUES (?, ?, ?, ?, ?, ?) ON CONFLICT(character_id) DO UPDATE SET name=excluded.name, ccu_level=excluded.ccu_level, ic_level=excluded.ic_level''', char)
    conn.commit()
    cursor.execute("SELECT character_id, name, ccu_level as ccu, ic_level as ic FROM characters")
    saved_chars = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return {"status": "success", "characters": saved_chars}

@app.get("/api/initial-data")
def get_initial_data():
    df = load_planet_data()
    bases = sorted(list(df['Constellation'].unique()))
    with open('recipes.json', 'r', encoding='utf-8') as f:
        recipes = json.load(f)
    products = [k for k, v in recipes.items() if v.get('type') in ['P2', 'P3', 'P4']]
    return {"status": "success", "bases": bases, "products": sorted(products), "product_ids": PI_TYPE_IDS}

@app.get("/api/market/best-product")
async def get_best_production():
    prices = await get_market_prices()
    if not prices: raise HTTPException(status_code=500, detail="Market data unavailable")
    profitability = [{"product": k, "jita_buy": v} for k, v in prices.items() if k not in ["Water", "Oxygen"]]
    profitability.sort(key=lambda x: x["jita_buy"], reverse=True)
    return {"status": "success", "recommended": profitability[0]["product"] if profitability else "Robotics", "analytics": profitability[:50]}

@app.post("/api/systems")
def get_systems(req: SystemRequest):
    df = load_planet_data()
    if not req.constellations: return {"status": "success", "systems": []}
    sys = sorted(list(df[df['Constellation'].isin(req.constellations)]['System'].dropna().unique()))
    return {"status": "success", "systems": sys}

@app.post("/api/calculate")
async def calculate_plan(req: PlanRequest):
    try:
        if not req.target_products: return {"status": "success", "data": [], "warning": "Выберите хотя бы один продукт."}
            
        with open('recipes.json', 'r', encoding='utf-8') as f: recipes = json.load(f)
            
        def build_chain(prod, requested_factories):
            rec = recipes.get(prod)
            if not rec: return None
            ptype = rec['type']
            node = {'prod': prod, 'type': ptype, 'factories': requested_factories, 'children': []}
            multiplier = 2 if ptype in ['P4', 'P3'] else 1
            for inp in rec.get('inputs', {}):
                child = build_chain(inp, requested_factories * multiplier)
                if child: node['children'].append(child)
            return node
            
        def aggregate_reqs(current_tree):
            reqs = {}
            def traverse(node):
                p = node['prod']
                if p not in reqs: reqs[p] = {'type': node['type'], 'factories': 0, 'inputs': recipes.get(p, {}).get('inputs', {})}
                reqs[p]['factories'] += node['factories']
                for c in node['children']: traverse(c)
            if current_tree: traverse(current_tree)
            return reqs

        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT character_id, name, ccu_level as ccu, ic_level as ic FROM characters")
        all_chars = [dict(row) for row in cursor.fetchall()]
        conn.close()
        
        f_chars = [c for c in all_chars if c['ccu'] == 5]
        m_chars = [c for c in all_chars if c['ccu'] < 5]
        f_slots = [c for c in f_chars for _ in range(c['ic'] + 1)]
        m_slots = [c for c in m_chars for _ in range(c['ic'] + 1)]
        total_f_planets, total_m_planets = len(f_slots), len(m_slots)
        
        req_f_base = req_m_base = 0
        for target in req.target_products:
            target_ptype = recipes.get(target, {}).get('type', 'P4')
            base_factories = 16 if target_ptype == 'P4' else 24
            reqs_base = aggregate_reqs(build_chain(target, base_factories))
            ht_base = sum(d['factories'] for d in reqs_base.values() if d['type'] == 'P4')
            adv_base = sum(d['factories'] for d in reqs_base.values() if d['type'] in ['P2', 'P3'])
            p1_base = sum(d['factories'] for d in reqs_base.values() if d['type'] == 'P1')
            req_f_base += math.ceil(ht_base / 16) + math.ceil(adv_base / 24)
            req_m_base += math.ceil(p1_base / 12)

        chain_multiplier = min(total_f_planets // req_f_base, (total_f_planets + total_m_planets) // (req_f_base + req_m_base)) if req_f_base > 0 and req_m_base > 0 else 1
        if chain_multiplier < 1: chain_multiplier = 1

        combined_prod_reqs = {}
        for target in req.target_products:
            target_ptype = recipes.get(target, {}).get('type', 'P4')
            base_factories = 16 if target_ptype == 'P4' else 24
            reqs = aggregate_reqs(build_chain(target, base_factories * chain_multiplier))
            for p, d in reqs.items():
                if p not in combined_prod_reqs: combined_prod_reqs[p] = {'type': d['type'], 'factories': 0, 'inputs': d['inputs']}
                combined_prod_reqs[p]['factories'] += d['factories']

        ht_facts = sum(d['factories'] for d in combined_prod_reqs.values() if d['type'] == 'P4')
        adv_facts = sum(d['factories'] for d in combined_prod_reqs.values() if d['type'] in ['P2', 'P3'])
        req_total_f_planets = (math.ceil(ht_facts / 16) if ht_facts else 0) + (math.ceil(adv_facts / 24) if adv_facts else 0)
        req_total_m_planets = sum(math.ceil(d['factories'] / 12) for d in combined_prod_reqs.values() if d['type'] == 'P1')
        leftover_f = max(0, total_f_planets - req_total_f_planets)
        available_m = total_m_planets + leftover_f

        warnings_list = []
        is_deficit = False
        if req_total_f_planets > total_f_planets:
            is_deficit = True
            warnings_list.append(f"ДЕФИЦИТ СБОРКИ: Требуется еще {req_total_f_planets - total_f_planets} планет под заводы")
        if req_total_m_planets > available_m:
            is_deficit = True
            warnings_list.append(f"ДЕФИЦИТ СЫРЬЯ: Требуется еще {req_total_m_planets - available_m} добывающих планет")

        is_excess = not is_deficit and (total_f_planets - req_total_f_planets >= 1) and (available_m - req_total_m_planets >= 2)
        warning = " | ⚠️ ".join(["⚠️ " + w for w in warnings_list]) if warnings_list else None

        # --- LOGISTICS ADVISOR LOGIC ---
        recommendation = None
        if (is_deficit or is_excess) and total_f_planets > 0:
            prices = await get_market_prices()
            if prices:
                fitting_products = []
                for prod, data in recipes.items():
                    if data.get('type') not in ['P2', 'P3', 'P4']: continue
                    t_ptype = data.get('type')
                    b_fact = 16 if t_ptype == 'P4' else 24
                    t_reqs = aggregate_reqs(build_chain(prod, b_fact))
                    t_ht = sum(d['factories'] for d in t_reqs.values() if d['type'] == 'P4')
                    t_adv = sum(d['factories'] for d in t_reqs.values() if d['type'] in ['P2', 'P3'])
                    t_rf = math.ceil(t_ht / 16) + math.ceil(t_adv / 24)
                    t_rm = sum(math.ceil(d['factories'] / 12) for d in t_reqs.values() if d['type'] == 'P1')
                    
                    if t_rf <= total_f_planets and t_rm <= total_m_planets + (total_f_planets - t_rf):
                        fitting_products.append(prod)
                
                if fitting_products:
                    fitting_products.sort(key=lambda x: prices.get(x, 0), reverse=True)
                    best_alt = fitting_products[0]
                    best_price = prices.get(best_alt, 0)
                    
                    if best_alt not in req.target_products:
                        if is_deficit:
                            recommendation = {
                                "type": "deficit", "product": best_alt,
                                "message_ru": f"Вам не хватает планет для текущего плана. Рекомендуем переключиться на «{best_alt}» — это самый дорогой товар, который идеально впишется в ваш лимит.",
                                "message_en": f"Capacity deficit. We recommend switching to «{best_alt}» — the most profitable product that fits perfectly within your planetary limits."
                            }
                        elif is_excess:
                            current_price = sum(prices.get(p, 0) for p in req.target_products)
                            if best_price > current_price * 1.15:
                                recommendation = {
                                    "type": "excess", "product": best_alt,
                                    "message_ru": f"У вас простаивают мощности. Вы можете производить «{best_alt}» — это принесет больше прибыли и задействует планеты эффективнее.",
                                    "message_en": f"You have idle capacity. You can produce «{best_alt}» for higher profit and better utilization of your planets."
                                }

        df = load_planet_data()
        search_df = df[df['Constellation'].isin(req.constellations)].copy() if req.constellations else df.copy()
        factory_const_series = df[df['System'] == req.factory_sys]['Constellation']
        factory_const = factory_const_series.iloc[0] if not factory_const_series.empty else "Unknown"
        
        plan, factory_tasks = [], []
        for p, d in combined_prod_reqs.items():
            if d['type'] in ['P2', 'P3', 'P4']:
                ins = ", ".join(d['inputs'].keys()) if d['inputs'] else "Компоненты"
                facts_left, chunk_size = d['factories'], 16 if d['type'] == 'P4' else 24
                while facts_left > 0:
                    allocate = min(facts_left, chunk_size)
                    factory_tasks.append({"prod": p, "type": d['type'], "facts": allocate, "ins": ins})
                    facts_left -= allocate
                
        factory_tasks.sort(key=lambda x: {'P2': 1, 'P3': 2, 'P4': 3}.get(x['type'], 0))
        for t in factory_tasks:
            char = f_slots.pop(0) if f_slots else {"name": "Нет альта", "ccu": 5, "character_id": 1}
            plan.append({
                "constellation": factory_const, "system": req.factory_sys, "planet": "Любая Barren",
                "character": char['name'], "char_id": char['character_id'],
                "assignment": f"Сборка {t['type']} ({t['ins']})", "res_out": t['prod'], 
                "type_id": PI_TYPE_IDS.get(t['prod'], 0), "structures": f"{t['facts']} заводов",
                "role": f"🏭 Переработка {t['type']} (x{chain_multiplier})", "res_in": t['ins'], 
                "cc_type": "1x Barren Command Center", "pg_load": 95, "cpu_load": 95
            })

        combined_m_slots, p0_needs, miner_tasks, missing_resources = m_slots + f_slots, {}, [], set()
        for p, d in combined_prod_reqs.items():
            if d['type'] == 'P1':
                p0 = recipes[p]['source']
                if p0 not in p0_needs: p0_needs[p0] = {'p1': p, 'density': 0, 'needed_planets': 0}
                p0_needs[p0]['needed_planets'] += math.ceil(d['factories'] / 12)
                
        for p0 in p0_needs:
            if p0 in search_df.columns:
                search_df[p0] = pd.to_numeric(search_df[p0], errors='coerce')
                avail = search_df.dropna(subset=[p0])
                if avail.empty: missing_resources.add(p0)
                else: p0_needs[p0]['density'] = avail[p0].max()
            else: missing_resources.add(p0)
            if p0 not in missing_resources: miner_tasks.extend([p0] * p0_needs[p0]['needed_planets'])
                
        excess = len(combined_m_slots) - len(miner_tasks)
        if excess > 0 and p0_needs and not missing_resources:
            valid_p0 = sorted([k for k in p0_needs.keys() if k not in missing_resources], key=lambda x: p0_needs[x]['density'])
            if valid_p0: miner_tasks.extend([valid_p0[i % len(valid_p0)] for i in range(excess)])
                
        if missing_resources:
            missing_str = ", ".join(missing_resources)
            warning = f"{warning} | ⚠️ КРИТИЧЕСКИЙ ДЕФИЦИТ: Нет планет для добычи: {missing_str}!" if warning else f"⚠️ КРИТИЧЕСКИЙ ДЕФИЦИТ: Нет планет для добычи: {missing_str}!"

        p0_alloc_counts = {p0: 0 for p0 in p0_needs}
        for p0 in miner_tasks:
            avail = search_df.dropna(subset=[p0])
            if not avail.empty:
                avail_sorted = avail.sort_values(by=p0, ascending=False)
                bp = avail_sorted.iloc[p0_alloc_counts[p0] % len(avail_sorted)]
                role_text = "⛏ Добыча (Избыток)" if p0_alloc_counts[p0] >= p0_needs[p0]['needed_planets'] else "⛏ Добыча P1"
                p0_alloc_counts[p0] += 1
                char = combined_m_slots.pop(0) if combined_m_slots else {"name": "Нет альта", "ccu": 0, "character_id": 1}
                out_p1 = p0_needs[p0]['p1']
                plan.append({
                    "constellation": bp['Constellation'], "system": bp['System'], "planet": f"Planet {bp['Planet']} ({bp['Type']})",
                    "character": char['name'], "char_id": char['character_id'], "assignment": f"Добыча P0 -> P1 ({p0})",
                    "res_out": out_p1, "type_id": PI_TYPE_IDS.get(out_p1, 0), "structures": "Экстрактор + 12 Basic Factories",
                    "role": role_text, "res_in": "-", "cc_type": f"1x {bp['Type']} Command Center", "pg_load": 85, "cpu_load": 75
                })
                
        return {"status": "success", "data": plan, "warning": warning, "recommendation": recommendation}
    except Exception as e:
        print(traceback.format_exc())
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/export")
def export_plan(req: ExportRequest):
    try:
        formatted_data = []
        cc_summary = {}
        for row in req.plan_data:
            cc_type = row.get('cc_type', '')
            if cc_type: cc_summary[cc_type.replace("1x ", "")] = cc_summary.get(cc_type.replace("1x ", ""), 0) + 1
            formatted_data.append({"Констелляция": row.get('constellation', ''), "Система": row.get('system', ''), "Планета": row.get('planet', ''), "Персонаж": row.get('character', ''), "Назначение": row.get('assignment', ''), "Выходной ресурс": row.get('res_out', ''), "Количество заводов": row.get('structures', ''), "Командный центр": cc_type})
        df_plan = pd.DataFrame(formatted_data)
        df_summary = pd.DataFrame([{"Тип Командного Центра": k, "Количество": v} for k, v in cc_summary.items()])
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df_plan.to_excel(writer, index=False, sheet_name='PI Plan')
            ws_plan = writer.sheets['PI Plan']
            fill_h = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
            for c in ws_plan[1]: c.fill, c.font = fill_h, Font(color="FFFFFF", bold=True)
            for col in ws_plan.columns: ws_plan.column_dimensions[col[0].column_letter].width = max((len(str(c.value)) for c in col if c.value), default=0) + 2
            if not df_summary.empty:
                df_summary.to_excel(writer, index=False, sheet_name='Shopping List')
                ws_summary = writer.sheets['Shopping List']
                fill_shop = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                for c in ws_summary[1]: c.fill, c.font = fill_shop, Font(color="FFFFFF", bold=True)
                for col in ws_summary.columns: ws_summary.column_dimensions[col[0].column_letter].width = max((len(str(c.value)) for c in col if c.value), default=0) + 2
        output.seek(0)
        return StreamingResponse(output, headers={'Content-Disposition': 'attachment; filename="PI_Plan.xlsx"'}, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

if __name__ == '__main__':
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)